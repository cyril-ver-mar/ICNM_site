#!/usr/bin/env python3
"""Build Word (.docx) templates for institute staff data collection.

Requires: python-docx (`pip install python-docx` in the project venv).
Optional screenshots: docs/work/colleague-packet/templates/_screens/*.png
  (regenerate with scripts/capture_colleague_screens.py).

Output: docs/work/colleague-packet/templates/*.docx
Regenerate: python scripts/build_colleague_templates.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "work" / "colleague-packet" / "templates"
SCREENS = OUT / "_screens"

PREVIEW_URL = "https://cyril-ver-mar.github.io/ICNM_site/preview-filled/"

LANG_WARN = (
    "Важно про языки (RU / EN / BE / ZH): заполняйте все четыре, если можете. "
    "Русский — обязательный минимум. Если язык не заполнен, текст будет "
    "машинно переведён с русского (или с единственного заполненного языка). "
    "Машинный перевод перед публикацией нужно проверить."
)

MULTI_AFFIL = (
    "Несколько ролей / подразделений: один человек может одновременно быть, "
    "например, сотрудником лаборатории и председателем профсоюза, "
    "заместителем директора и учёным секретарём. В шаблоне персоналий "
    "укажите ВСЕ роли и подразделения (affiliations). Тогда на персональной "
    "странице people/{id}/ и в пакетах подразделений отобразятся все связи. "
    "Не заводите дубликаты одного человека под разными id."
)

PHOTO_NOTE = (
    "Фото: отдельным файлом JPG/PNG. Имя файла — как в поле «Имя файла фото». "
    "Верните фото вместе с заполненным Word-файлом (письмо разработчику "
    "или папка assets/incoming/rosters/photos/)."
)


def _set_run_font(run, *, bold: bool = False, size: int = 11, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    if color is not None:
        run.font.color.rgb = color


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def _add_para(
    doc: Document,
    text: str,
    *,
    bold: bool = False,
    size: int = 11,
    space_after: int = 6,
    color: RGBColor | None = None,
) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, bold=bold, size=size, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)


def _add_warn_box(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, bold=True, size=11, color=RGBColor(0x66, 0x4D, 0x03))
    p.paragraph_format.space_after = Pt(10)
    # Light yellow shading on paragraph
    shd = p._p.get_or_add_pPr()
    from docx.oxml import OxmlElement

    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "FFF3CD")
    shade.set(qn("w:val"), "clear")
    shd.append(shade)


def _add_screenshot(doc: Document, name: str, caption: str) -> None:
    path = SCREENS / f"{name}.png"
    _add_para(doc, caption, bold=True, size=11, space_after=4)
    if path.is_file():
        doc.add_picture(str(path), width=Cm(16.5))
        cap = doc.add_paragraph()
        run = cap.add_run(f"Скриншот наполненного макета ({name}). Это рыба для ориентира, не официальные данные.")
        _set_run_font(run, size=9, color=RGBColor(0x55, 0x55, 0x55))
        cap.paragraph_format.space_after = Pt(10)
    else:
        _add_para(
            doc,
            f"[Нет файла _screens/{name}.png — запустите scripts/capture_colleague_screens.py]",
            size=10,
            color=RGBColor(0x99, 0x00, 0x00),
        )


def _add_field_table(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    """rows: (field_key, russian_label, example_or_hint)"""
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, title in enumerate(("Поле (ключ)", "Что заполнять", "Пример / подсказка")):
        hdr[i].text = title
        for p in hdr[i].paragraphs:
            for run in p.runs:
                _set_run_font(run, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF))
            # navy header fill
            from docx.oxml import OxmlElement

            tc = hdr[i]._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "0F2A43")
            shd.set(qn("w:val"), "clear")
            tcPr.append(shd)

    for key, label, hint in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = label
        cells[2].text = hint
        for cell in cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    _set_run_font(run, size=10)

    doc.add_paragraph()


def _add_blank_entry_block(doc: Document, title: str, fields: list[str], count: int = 3) -> None:
    _add_heading(doc, title, level=2)
    _add_para(
        doc,
        "Заполните блоки ниже (скопируйте блок, если записей больше). "
        "Пустые языковые поля оставьте пустыми — сработает машинный перевод.",
        size=10,
    )
    for n in range(1, count + 1):
        _add_para(doc, f"— Запись {n} —", bold=True, size=11, space_after=4)
        table = doc.add_table(rows=len(fields), cols=2)
        table.style = "Table Grid"
        for i, field in enumerate(fields):
            table.rows[i].cells[0].text = field
            table.rows[i].cells[1].text = ""
            for p in table.rows[i].cells[0].paragraphs:
                for run in p.runs:
                    _set_run_font(run, bold=True, size=10)
            for p in table.rows[i].cells[1].paragraphs:
                for run in p.runs:
                    _set_run_font(run, size=10)
        doc.add_paragraph()


def _common_front(
    doc: Document,
    title: str,
    purpose: list[str],
    screens: list[tuple[str, str]],
) -> None:
    _add_heading(doc, title, level=0)
    _add_para(
        doc,
        "Памятка для сотрудников ИХНМ. Заполненный файл верните разработчику сайта "
        "(письмо) или положите в assets/incoming/rosters/.",
        size=11,
    )
    _add_para(doc, f"Живой макет сайта (рыба): {PREVIEW_URL}", size=11)
    _add_warn_box(doc, LANG_WARN)
    _add_warn_box(doc, MULTI_AFFIL)
    _add_para(doc, PHOTO_NOTE, size=10)

    _add_heading(doc, "Зачем и куда на сайте", level=1)
    for line in purpose:
        _add_para(doc, "• " + line, size=11, space_after=4)

    _add_heading(doc, "Как это выглядит на сайте (наполненный макет)", level=1)
    _add_para(
        doc,
        "Ниже — скриншоты страниц с «рыбой». По ним видно, куда попадёт каждая строка. "
        "Реальные тексты и фото подставят после ваших файлов.",
        size=11,
    )
    for name, caption in screens:
        _add_screenshot(doc, name, caption)


def _save(doc: Document, name: str) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    doc.save(path)
    return path


# --- Field catalogues (Russian labels) ---

PEOPLE_FIELDS = [
    ("id", "Короткий латинский идентификатор (без пробелов)", "ivanova, rogachev"),
    ("name_ru", "ФИО полностью — русский", "Иванова Наталья Александровна"),
    ("name_en", "Full name — English", "Natalia Ivanova"),
    ("name_be", "ПІБ — беларуская", ""),
    ("name_zh", "姓名 — 中文", ""),
    ("role_ru", "Основная должность / роль на карточке — русский", "Заведующий лабораторией"),
    ("role_en", "Role — English", "Head of laboratory"),
    ("role_be", "Пасада — беларуская", ""),
    ("role_zh", "职务 — 中文", ""),
    ("degree_ru", "Учёная степень / звание — русский", "к.х.н."),
    ("degree_en", "Degree — English", "PhD"),
    ("degree_be", "Вучоная ступень — беларуская", ""),
    ("degree_zh", "学位 — 中文", ""),
    ("phone", "Телефон", "+375 (17) 000-00-00"),
    ("email", "Электронная почта", "ivanova@ichnm.by"),
    (
        "affiliations",
        "ВСЕ подразделения и роли: unit_id|роль; несколько через ;",
        "lab-films|Заведующий лабораторией; union|Председатель профсоюза",
    ),
    ("bio_ru", "Краткая биография / научный профиль — русский", ""),
    ("bio_en", "Bio — English", ""),
    ("bio_be", "Біяграфія — беларуская", ""),
    ("bio_zh", "简介 — 中文", ""),
    ("orcid_url", "Ссылка ORCID (если есть метрики — URL обязателен)", "https://orcid.org/…"),
    ("google_scholar_url", "Ссылка Google Scholar", ""),
    ("scopus_author_url", "Ссылка Scopus Author", ""),
    ("elibrary_url", "Ссылка eLIBRARY / РИНЦ", ""),
    ("researchgate_url", "Ссылка ResearchGate", ""),
    ("h_index_google_scholar", "h-индекс Google Scholar (число)", "12"),
    ("citations_google_scholar", "Цитирования Google Scholar (число)", "340"),
    ("h_index_scopus", "h-индекс Scopus", ""),
    ("citations_scopus", "Цитирования Scopus", ""),
    ("awards_ru", "Награды (по строке; только если есть)", ""),
    ("publications_scientific_ru", "Свои научные публикации на странице (опционально)", ""),
    ("publications_methodical_ru", "Методические публикации (опционально)", ""),
    ("interests_ru", "Научные интересы (опционально)", ""),
    ("projects_ru", "Проекты на персональной странице (опционально)", ""),
    ("photo_filename", "Имя файла фото", "ivanova.jpg"),
    ("notes", "Заметки для разработчика (не на сайт)", ""),
]

LAB_FIELDS = [
    ("slug", "Код лаборатории", "nano | films | lcd | composites | woodchem"),
    ("title_ru", "Название лаборатории — русский", ""),
    ("title_en", "Lab title — English", ""),
    ("title_be", "Назва — беларуская", ""),
    ("title_zh", "名称 — 中文", ""),
    ("kicker_ru", "Короткая подпись под названием — русский", "Подразделение Института"),
    ("about_ru", "Текст «О лаборатории» — русский", ""),
    ("about_en", "About — English", ""),
    ("about_be", "Пра лабараторыю — беларуская", ""),
    ("about_zh", "简介 — 中文", ""),
    ("head_id", "id персоналии заведующего (из 01-personalia)", "kulikouskaya"),
    ("phone", "Телефон лаборатории", ""),
    ("email", "E-mail лаборатории", ""),
    ("notes", "Заметки", ""),
]

LAB_PROJECT_FIELDS = [
    ("lab_slug", "Код лаборатории", "films"),
    ("title_ru", "Название проекта — русский", ""),
    ("title_en", "Project title — English", ""),
    ("title_be", "Назва праекта — беларуская", ""),
    ("title_zh", "项目名称 — 中文", ""),
    ("years", "Годы", "2024–2026"),
    ("status", "Статус: active или completed", "active"),
    ("lead_ru", "Краткое описание — русский", ""),
    ("lead_en", "Short description — English", ""),
]

DEV_FIELDS = [
    ("slug", "Латинский slug страницы", "polaroid-films"),
    ("title_ru", "Название разработки — русский", ""),
    ("title_en", "Title — English", ""),
    ("title_be", "Назва — беларуская", ""),
    ("title_zh", "名称 — 中文", ""),
    ("lead_ru", "Краткий лид для плитки — русский", ""),
    ("lead_en", "Tile lead — English", ""),
    ("lead_be", "Лід — беларуская", ""),
    ("lead_zh", "摘要 — 中文", ""),
    ("description_ru", "Развёрнутое описание — русский", ""),
    ("product_or_spec_ru", "Продукт / теххарактеристики — русский", ""),
    ("lab_slug", "Лаборатория-владелец", "films"),
    ("staff_id", "id закреплённого сотрудника", "ivanova"),
    ("direction_slug", "Slug связанного направления (если есть)", "thin-films"),
    ("photo_filename", "Имя файла фото", "polaroid.jpg"),
    ("notes", "Заметки", ""),
]

DIR_FIELDS = [
    ("slug", "Латинский slug", "thin-films"),
    ("title_ru", "Название направления — русский", ""),
    ("title_en", "Title — English", ""),
    ("title_be", "Назва — беларуская", ""),
    ("title_zh", "名称 — 中文", ""),
    ("lead_ru", "Краткий лид — русский", ""),
    ("lead_en", "Lead — English", ""),
    ("lead_be", "Лід — беларуская", ""),
    ("lead_zh", "摘要 — 中文", ""),
    ("description_ru", "Описание — русский", ""),
    ("scope", "institute (общеинститутское) или lab", "lab"),
    ("lab_slug", "Если scope=lab — код лаборатории", "films"),
    ("staff_id", "id ответственного (опционально)", "ivanova"),
    ("notes", "Заметки", ""),
]

HR_UNIT_FIELDS = [
    ("unit_id", "Код подразделения (не менять)", "hr"),
    ("title_ru", "Название — русский", "Отдел кадров"),
    ("title_en", "Title — English", "HR department"),
    ("phone", "Телефон отдела", "+375 (17) 243-67-56"),
    ("email", "E-mail отдела", ""),
    ("about_ru", "Кратко о функциях — русский", ""),
]

HR_STAFF_FIELDS = [
    ("person_id", "id из 01-personalia", "hr-head"),
    ("role_ru", "Роль в отделе — русский", "Начальник отдела кадров"),
    ("role_en", "Role — English", ""),
    ("role_be", "Роля — беларуская", ""),
    ("role_zh", "职务 — 中文", ""),
    ("phone", "Телефон (если отличается)", ""),
    ("email", "E-mail", ""),
    ("notes", "Заметки", ""),
]

ENG_OT_FIELDS = [
    ("unit_id", "Код: engineering или labor-protection", "engineering"),
    ("unit_title_ru", "Название подразделения — русский", "Главный инженер"),
    ("phone", "Телефон подразделения", ""),
    ("person_id", "id персоналии", "tikhonov"),
    ("role_ru", "Роль — русский", "Главный инженер"),
    ("role_en", "Role — English", ""),
    ("role_be", "Роля — беларуская", ""),
    ("role_zh", "职务 — 中文", ""),
    ("phone_person", "Телефон сотрудника", ""),
    ("email", "E-mail", ""),
    ("notes", "Заметки", ""),
]

UNION_FIELDS = [
    ("person_id", "id из 01-personalia (предпочтительно)", "union-chair"),
    ("name_ru", "ФИО — русский (если ещё нет в персоналиях)", ""),
    ("name_en", "Name — English", ""),
    ("name_be", "ПІБ — беларуская", ""),
    ("name_zh", "姓名 — 中文", ""),
    ("role_ru", "Роль в профсоюзе — русский", "Председатель профсоюза"),
    ("role_en", "Role — English", ""),
    ("role_be", "Роля — беларуская", ""),
    ("role_zh", "职务 — 中文", ""),
    ("phone", "Телефон", ""),
    ("email", "E-mail", ""),
    ("photo_filename", "Имя файла фото", ""),
    ("notes", "Заметки", ""),
]


def build_people() -> Path:
    doc = Document()
    _common_front(
        doc,
        "01 — Персоналии (people/{id}/)",
        [
            "Карточка сотрудника на сайте и в составах лабораторий / подразделений.",
            "Уйдёт в CPT «Персоналии», страница people/{id}/; карточки в руководстве, лабораториях, кадрах, профсоюзе.",
            "На странице: фото, должность, контакты, все подразделения, наукометрия, опциональные блоки (награды, публикации, интересы, проекты) — только если заполнены.",
            "Сначала этот файл, потом составы в шаблонах лабораторий / кадров / профсоюза ссылаются на id.",
        ],
        [
            ("person-full", "Персональная страница (полностью заполненный пример)"),
            ("structure", "Структура — плитки подразделений, куда ведут карточки людей"),
            ("union", "Профсоюз — пример второй роли того же человека"),
        ],
    )
    _add_heading(doc, "Справочник полей", level=1)
    _add_field_table(doc, PEOPLE_FIELDS)
    entry_labels = [f"{k} — {label}" for k, label, _ in PEOPLE_FIELDS]
    _add_blank_entry_block(doc, "Данные для заполнения", entry_labels, count=4)
    return _save(doc, "01-personalia.docx")


def build_labs() -> Path:
    doc = Document()
    _common_front(
        doc,
        "02 — Лаборатории (labs/{slug}/)",
        [
            "Пакет страницы лаборатории: о лаборатории, направления, проекты, оборудование, разработки, команда, публикации, контакты.",
            "Уйдёт на labs/{slug}/ и в пункт «Структура». Каталоги направлений / разработок / приборов подтягиваются по lab_slug.",
            "Пустые блоки на сайте скрываются.",
            "Состав команды — через персоналии (affiliations на lab-*). Проекты — отдельным блоком ниже.",
        ],
        [
            ("lab-pack", "Пакет лаборатории (рыба): меню разделов, направления, проекты, оборудование"),
            ("structure", "Плитка лаборатории на странице «Структура»"),
        ],
    )
    _add_heading(doc, "Справочник полей лаборатории", level=1)
    _add_field_table(doc, LAB_FIELDS)
    _add_blank_entry_block(
        doc,
        "Данные лабораторий",
        [f"{k} — {label}" for k, label, _ in LAB_FIELDS],
        count=5,
    )
    _add_heading(doc, "Проекты (действующие и завершённые)", level=1)
    _add_para(
        doc,
        "Список остаётся на странице лаборатории (не отдельный пункт верхнего меню). "
        "status: active = действующий, completed = завершённый.",
        size=11,
    )
    _add_field_table(doc, LAB_PROJECT_FIELDS)
    _add_blank_entry_block(
        doc,
        "Данные проектов",
        [f"{k} — {label}" for k, label, _ in LAB_PROJECT_FIELDS],
        count=4,
    )
    return _save(doc, "02-laboratories.docx")


def build_developments() -> Path:
    doc = Document()
    _common_front(
        doc,
        "03 — Разработки (CPT development)",
        [
            "Карточка разработки в институтском каталоге и в пакете лаборатории («Услуги и разработки»).",
            "Уйдёт в CPT «Разработки», developments/{slug}/, плитки на /developments/ и в labs/{slug}/.",
            "На странице: узкое фото, описание, теххарактеристики, лаборатория и сотрудник, связь с направлением.",
        ],
        [
            ("development", "Страница разработки (рыба)"),
            ("lab-pack", "Где разработки видны в пакете лаборатории"),
        ],
    )
    _add_heading(doc, "Справочник полей", level=1)
    _add_field_table(doc, DEV_FIELDS)
    _add_blank_entry_block(
        doc,
        "Данные разработок",
        [f"{k} — {label}" for k, label, _ in DEV_FIELDS],
        count=4,
    )
    return _save(doc, "03-developments.docx")


def build_directions() -> Path:
    doc = Document()
    _common_front(
        doc,
        "04 — Направления работы (CPT direction)",
        [
            "Позиции каталога «Направления работы» и блоки направлений в пакете лаборатории.",
            "Уйдёт в CPT «Направления», science/{slug}/, плитки на /science/, раздел #directions лаборатории.",
            "На детальной странице — связанные разработки.",
        ],
        [
            ("direction", "Страница направления (рыба)"),
            ("lab-pack", "Направления внутри пакета лаборатории"),
        ],
    )
    _add_heading(doc, "Справочник полей", level=1)
    _add_field_table(doc, DIR_FIELDS)
    _add_blank_entry_block(
        doc,
        "Данные направлений",
        [f"{k} — {label}" for k, label, _ in DIR_FIELDS],
        count=4,
    )
    return _save(doc, "04-science-directions.docx")


def build_hr() -> Path:
    doc = Document()
    _common_front(
        doc,
        "05 — Отдел кадров",
        [
            "Страница подразделения «Отдел кадров» в Структуре и контакты для справок «Для сотрудника».",
            "Уйдёт на /hr/, плитки сотрудников → people/{id}/.",
            "Сначала заведите людей в 01-personalia.docx (affiliations: hr|…).",
        ],
        [
            ("structure", "Плитки административных подразделений на «Структуре»"),
            ("person-full", "Куда ведёт плитка сотрудника"),
        ],
    )
    _add_heading(doc, "Подразделение", level=1)
    _add_field_table(doc, HR_UNIT_FIELDS)
    _add_blank_entry_block(
        doc,
        "Реквизиты отдела (одна запись)",
        [f"{k} — {label}" for k, label, _ in HR_UNIT_FIELDS],
        count=1,
    )
    _add_heading(doc, "Состав", level=1)
    _add_field_table(doc, HR_STAFF_FIELDS)
    _add_blank_entry_block(
        doc,
        "Сотрудники отдела кадров",
        [f"{k} — {label}" for k, label, _ in HR_STAFF_FIELDS],
        count=4,
    )
    return _save(doc, "05-hr.docx")


def build_engineering_ot() -> Path:
    doc = Document()
    _common_front(
        doc,
        "06 — Главный инженер и охрана труда",
        [
            "Две admin-unit страницы в Структуре: «Главный инженер» (engineering) и «Охрана труда» (labor-protection).",
            "Уйдёт на /engineering/, /labor-protection/, карточки → people/{id}/.",
            "Персоналии заведите в 01-personalia.docx.",
        ],
        [
            ("structure", "Плитки на «Структуре»"),
            ("person-full", "Персональная страница сотрудника"),
        ],
    )
    _add_heading(doc, "Справочник полей", level=1)
    _add_field_table(doc, ENG_OT_FIELDS)
    _add_blank_entry_block(
        doc,
        "Главный инженер — состав",
        [f"{k} — {label}" for k, label, _ in ENG_OT_FIELDS],
        count=3,
    )
    _add_blank_entry_block(
        doc,
        "Охрана труда — состав",
        [f"{k} — {label}" for k, label, _ in ENG_OT_FIELDS],
        count=3,
    )
    return _save(doc, "06-engineering-labor-protection.docx")


def build_union() -> Path:
    doc = Document()
    _common_front(
        doc,
        "07 — Профсоюз (первичная организация)",
        [
            "Страница первичной профсоюзной организации Института (не замена profnan.by).",
            "Уйдёт на /union/, плитки → people/{id}/; ссылка на коллективный договор в Документах.",
            "В персоналиях укажите affiliations: union|Председатель профсоюза и т.п. "
            "Если человек ещё и в лаборатории — обе роли в одном id.",
        ],
        [
            ("union", "Страница профсоюза (рыба)"),
            ("structure", "Плитка профсоюза в Структуре (после admin-блока)"),
            ("person-full", "Персоналия с несколькими ролями"),
        ],
    )
    _add_heading(doc, "Текст страницы", level=1)
    _add_field_table(
        doc,
        [
            ("intro_ru", "Вводный текст — русский", "Первичная профсоюзная организация Института…"),
            ("intro_en", "Intro — English", ""),
            ("intro_be", "Уводзіны — беларуская", ""),
            ("intro_zh", "简介 — 中文", ""),
            ("plan_note_ru", "План работы / контакты бюро — русский", ""),
        ],
    )
    _add_blank_entry_block(
        doc,
        "Текст страницы (одна запись)",
        [
            "intro_ru — Вводный текст — русский",
            "intro_en — Intro — English",
            "intro_be — Уводзіны — беларуская",
            "intro_zh — 简介 — 中文",
            "plan_note_ru — План работы / контакты бюро — русский",
        ],
        count=1,
    )
    _add_heading(doc, "Состав профкома", level=1)
    _add_field_table(doc, UNION_FIELDS)
    _add_blank_entry_block(
        doc,
        "Члены / председатель",
        [f"{k} — {label}" for k, label, _ in UNION_FIELDS],
        count=4,
    )
    return _save(doc, "07-union.docx")


def main() -> None:
    # Remove legacy Excel templates if present (Word is primary).
    for old in OUT.glob("*.xlsx"):
        old.unlink()
        print(f"Removed legacy {old.relative_to(ROOT)}")

    paths = [
        build_people(),
        build_labs(),
        build_developments(),
        build_directions(),
        build_hr(),
        build_engineering_ot(),
        build_union(),
    ]
    for path in paths:
        print(f"Wrote {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
