#!/usr/bin/env python3
"""Copy docs/Mail.docx → docs/work/colleague-packet/Mail-rasylka.docx and append handoff notes.

Does NOT modify docs/Mail.docx.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "Mail.docx"
DST = ROOT / "docs" / "work" / "colleague-packet" / "Mail-rasylka.docx"

PREVIEW = "https://cyril-ver-mar.github.io/ICNM_site/preview-filled/"
LANDING = "https://cyril-ver-mar.github.io/ICNM_site/"
TEMPLATES = "docs/work/colleague-packet/templates/"


def _run(p, text: str, *, bold: bool = False, size: int = 12) -> None:
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _para(doc: Document, text: str, *, bold: bool = False, size: int = 12) -> None:
    p = doc.add_paragraph()
    _run(p, text, bold=bold, size=size)
    p.paragraph_format.space_after = Pt(8)


def main() -> None:
    if not SRC.is_file():
        raise SystemExit(f"Missing source letter: {SRC}")
    DST.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, DST)

    doc = Document(str(DST))
    doc.add_paragraph()
    _para(doc, "— — —", size=12)
    _para(
        doc,
        "Дополнение к письму: как пользоваться материалами",
        bold=True,
        size=14,
    )

    _para(doc, "1. Живой макет сайта (наполненный, с «рыбой»)", bold=True)
    _para(
        doc,
        f"Откройте в браузере: {PREVIEW}\n"
        f"Стартовая страница репозитория: {LANDING}\n"
        "Это черновик вида и структуры, не официальный ichnm.by. По макету видно, "
        "куда лягут ФИО, фото, роли, тексты лабораторий, направления и разработки.",
    )

    _para(doc, "2. Шаблоны-памятки (Word)", bold=True)
    _para(
        doc,
        f"В папке {TEMPLATES} семь файлов .docx:\n"
        "01-personalia.docx — персоналии (личные страницы);\n"
        "02-laboratories.docx — лаборатории и научные проекты;\n"
        "03-developments.docx — разработки;\n"
        "04-science-directions.docx — направления работы;\n"
        "05-hr.docx — отдел кадров;\n"
        "06-engineering-labor-protection.docx — главный инженер и охрана труда;\n"
        "07-union.docx — профсоюз.\n"
        "В начале каждой памятки — скриншоты из макета и жёлтые предупреждения про языки "
        "и про несколько ролей одного человека. Дальше — справочник полей на русском "
        "и пустые блоки для заполнения.",
    )

    _para(doc, "3. Языки RU / EN / BE / ZH", bold=True)
    _para(
        doc,
        "Структура сайта уже на четырёх языках. Тексты института желательно дать на всех. "
        "Русский — обязательный минимум. Пустой язык будет машинно переведён; "
        "китайский мы переведём машинно и отдадим носителю на проверку. "
        "EN/BE тоже можно оставить на машинный перевод, если есть только один язык.",
    )

    _para(doc, "4. Несколько ролей у одного сотрудника", bold=True)
    _para(
        doc,
        "Если человек, например, сотрудник лаборатории и одновременно председатель профсоюза "
        "(или заместитель директора и учёный секретарь) — в 01-personalia укажите ВСЕ "
        "подразделения и роли в одном id. Тогда сайт покажет их на персональной странице "
        "и в составах подразделений. Не создавайте дубликаты одного человека.",
    )

    _para(doc, "5. Фото", bold=True)
    _para(
        doc,
        "Отдельными файлами JPG/PNG. Имя файла — как в поле «Имя файла фото» "
        "(латиницей, без пробелов: ivanova.jpg). Портрет желательно вертикальный, "
        "лицо крупно. Приложите фото к письму вместе с Word-файлами.",
    )

    _para(doc, "6. Куда вернуть заполненное", bold=True)
    _para(
        doc,
        "Ответом на это письмо (вложите .docx и фото) разработчику сайта. "
        "Либо положите копии в assets/incoming/rosters/ (фото — в photos/). "
        "PDF устава, антикоррупции, эмблема НАН, URL соцсетей ИХНМ — отдельными "
        "файлами в assets/incoming/ (см. docs/work/notes/colleague-packet.md).",
    )

    _para(doc, "7. Порядок заполнения (рекомендуется)", bold=True)
    _para(
        doc,
        "Сначала 01-personalia (все люди и все их роли).\n"
        "Затем 02-laboratories, 03-developments, 04-science-directions.\n"
        "Потом 05–07 (кадры, инженер/ОТ, профсоюз) — там ссылки на id из персоналий.\n"
        "В текстах разработок и направлений ориентируйтесь на внешних партнёров и заказчиков.",
    )

    _para(doc, "8. Что пока не нужно в этих шаблонах", bold=True)
    _para(
        doc,
        "Новости, «СМИ о нас», мероприятия и таблица публикаций идут отдельно "
        "(после запуска их ведёт редактор лент; стартовый список публикаций — "
        "отдельным файлом в assets/incoming/publications/).",
    )

    doc.save(str(DST))
    print(f"Wrote {DST.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
